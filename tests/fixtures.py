"""Deterministic DOCX template fixtures and result-reading helpers.

Templates are built programmatically (not committed binaries) so the suite is
hermetic and nothing nondeterministic is churned into git. The committed on-disk
template under ``test/sql/data/`` is the one exception (the SQL E2E suite needs a
path to ATTACH against) and is generated once by ``make_template_file``.
"""

from __future__ import annotations

import io

from docx import Document
from docxtpl import DocxTemplate  # noqa: F401 -- imported so deps are exercised

# Known render values the tests assert on.
KNOWN_CUSTOMER = "Ada Lovelace"
KNOWN_TOTAL = "99.50"
KNOWN_ITEMS = [
    {"name": "Analytical Engine", "price": "1000"},
    {"name": "Difference Engine", "price": "500"},
]

# Jinja2/docxtpl template body: a placeholder header + a {% for %} loop.
_TEMPLATE_PARAS = [
    "Invoice for {{ customer }}",
    "Total: {{ total }}",
    "{% for item in items %}- {{ item.name }}: {{ item.price }}{% endfor %}",
]


def make_template_bytes() -> bytes:
    """A small .docx template with a {{ placeholder }} and a {% for %} loop."""
    doc = Document()
    for para in _TEMPLATE_PARAS:
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_garbage_bytes() -> bytes:
    """Not a .docx (no ZIP magic) -- must degrade to NULL / clean error."""
    return b"this is definitely not a docx file"


def docx_text(docx_bytes: bytes) -> str:
    """Extract concatenated paragraph text from rendered .docx bytes."""
    doc = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def make_template_file(path: str) -> None:
    """Write the standard template to ``path`` (for the SQL E2E suite)."""
    with open(path, "wb") as fh:
        fh.write(make_template_bytes())
