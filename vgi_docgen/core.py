"""Pure document-generation logic: DOCX templating + optional PDF conversion.

No Arrow, no VGI here -- just bytes in, bytes out, so this module is directly
unit-testable. The two public entry points are:

- :func:`render_docx` -- fill ONE DOCX template (via ``docxtpl`` / Jinja2-style
  ``{{ placeholders }}`` and ``{% for %}`` loops) with a single context dict,
  returning the rendered ``.docx`` bytes.
- :func:`merge_docx` -- fill ONE template with MANY contexts and concatenate the
  results into a single ``.docx`` (a page break between documents), for the
  "many SQL rows -> one merged document" path.

Templates resolve from one of two sources, in priority order:

1. A ``VARCHAR`` reference. If it looks like a path that exists (optionally under
   a configured templates directory, ``VGI_DOCGEN_TEMPLATES``), the file is read.
2. A ``BLOB`` of raw ``.docx`` template bytes passed inline.

PDF output is OPT-IN and degrades gracefully: it shells out to a headless
LibreOffice (``soffice``) if one is on PATH, and raises a clear
:class:`DocgenError` otherwise. LibreOffice is **never** a hard dependency.

Untrusted-binary discipline: every failure mode (missing template, malformed
DOCX, bad placeholder, Jinja error, missing LibreOffice) is funnelled into a
:class:`DocgenError` with an explicit message -- never an unhandled crash that
would take the worker down.
"""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docxcompose.composer import Composer
from docxtpl import DocxTemplate

__all__ = [
    "DocgenError",
    "TemplateRef",
    "render_docx",
    "merge_docx",
    "to_pdf",
    "libreoffice_available",
]

# Environment variable naming a directory that VARCHAR template refs resolve
# against (so callers can pass a bare ``invoice.docx`` instead of an abspath).
_TEMPLATES_ENV = "VGI_DOCGEN_TEMPLATES"

# The PK\x03\x04 ZIP local-file-header magic; every .docx is a ZIP container.
_ZIP_MAGIC = b"PK\x03\x04"


class DocgenError(Exception):
    """A clean, user-facing failure (bad template, render error, no LibreOffice).

    The worker turns these into NULL (scalar path) or a DuckDB error (merge
    path) -- never a crash.
    """


class TemplateRef:
    """A resolved template source: either a filesystem path or inline bytes.

    Construct via :meth:`from_path` (VARCHAR ref) or :meth:`from_bytes` (BLOB);
    both return ``None`` for a ``None`` input so NULLs pass straight through.
    """

    __slots__ = ("_data",)

    def __init__(self, data: bytes) -> None:
        self._data = data

    @classmethod
    def from_path(cls, ref: str | None, *, templates_dir: str | None = None) -> TemplateRef | None:
        """Resolve a VARCHAR template reference to bytes, or ``None`` for NULL.

        Looks for ``ref`` directly, then under ``templates_dir`` (falling back to
        ``$VGI_DOCGEN_TEMPLATES``). Raises :class:`DocgenError` if not found or
        not a readable ``.docx`` ZIP.
        """
        if ref is None:
            return None
        base = templates_dir or os.environ.get(_TEMPLATES_ENV)
        candidates: list[Path] = [Path(ref)]
        if base:
            candidates.append(Path(base) / ref)
        for candidate in candidates:
            if candidate.is_file():
                data = candidate.read_bytes()
                _check_docx_magic(data, source=str(candidate))
                return cls(data)
        raise DocgenError(f"template not found: {ref!r}")

    @classmethod
    def from_bytes(cls, ref: bytes | None) -> TemplateRef | None:
        """Wrap inline DOCX template bytes, or ``None`` for NULL.

        Raises :class:`DocgenError` if the bytes are not a ``.docx`` ZIP.
        """
        if ref is None:
            return None
        _check_docx_magic(ref, source="<inline bytes>")
        return cls(bytes(ref))

    @property
    def data(self) -> bytes:
        """The raw ``.docx`` template bytes."""
        return self._data


def _check_docx_magic(data: bytes, *, source: str) -> None:
    """Reject anything that is not a ZIP container before docxtpl touches it."""
    if len(data) < 4 or data[:4] != _ZIP_MAGIC:
        raise DocgenError(f"not a valid .docx (DOCX/ZIP) template: {source}")


def _load_template(ref: TemplateRef) -> DocxTemplate:
    """Build a fresh :class:`DocxTemplate` from template bytes.

    A new instance per render -- ``docxtpl`` mutates its document in place, so a
    template must never be shared across rows.
    """
    try:
        return DocxTemplate(io.BytesIO(ref.data))
    except Exception as exc:  # noqa: BLE001 -- untrusted binary; convert to clean error
        raise DocgenError(f"could not open DOCX template: {exc}") from exc


def render_docx(ref: TemplateRef, context: dict[str, Any]) -> bytes:
    """Render ONE document from ``ref`` filled with ``context``; return DOCX bytes.

    ``context`` keys become Jinja2 variables in the template
    (``{{ key }}``); nested dicts/lists drive ``{% for %}`` loops and tables.

    Raises:
        DocgenError: on a malformed template or a Jinja render error (bad
            placeholder, undefined behaviour, etc.).
    """
    tpl = _load_template(ref)
    try:
        tpl.render(context)
    except Exception as exc:  # noqa: BLE001 -- Jinja/template errors -> clean error
        raise DocgenError(f"template render failed: {exc}") from exc
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def merge_docx(ref: TemplateRef, contexts: list[dict[str, Any]]) -> bytes:
    """Render ``ref`` once per context and concatenate into ONE DOCX.

    Each rendered document is appended after a page break, producing a single
    merged ``.docx`` (the mail-merge "many rows -> one document" path). An empty
    ``contexts`` yields the template rendered against an empty context (a single
    empty document), so the output is always a valid ``.docx``.

    Raises:
        DocgenError: if any individual render fails, or the merge cannot
            assemble the documents.
    """
    if not contexts:
        return render_docx(ref, {})

    rendered = [render_docx(ref, ctx) for ctx in contexts]
    try:
        master = Document(io.BytesIO(rendered[0]))
        composer = Composer(master)
        for doc_bytes in rendered[1:]:
            master.add_page_break()
            composer.append(Document(io.BytesIO(doc_bytes)))
        out = io.BytesIO()
        composer.save(out)
        return out.getvalue()
    except Exception as exc:  # noqa: BLE001 -- merge plumbing -> clean error
        raise DocgenError(f"document merge failed: {exc}") from exc


def libreoffice_available() -> bool:
    """True if a headless LibreOffice (``soffice``) binary is on PATH."""
    return _soffice_path() is not None


def _soffice_path() -> str | None:
    """Locate a headless LibreOffice binary, honouring ``$VGI_DOCGEN_SOFFICE``."""
    override = os.environ.get("VGI_DOCGEN_SOFFICE")
    if override:
        return override if Path(override).exists() else shutil.which(override)
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return None


def to_pdf(docx_bytes: bytes) -> bytes:
    """Convert rendered DOCX bytes to PDF via headless LibreOffice.

    This is the OPT-IN PDF path. LibreOffice is an optional *runtime* dependency,
    never installed by this package.

    Raises:
        DocgenError: if no LibreOffice is found, or the conversion fails / times
            out / produces no PDF.
    """
    soffice = _soffice_path()
    if soffice is None:
        raise DocgenError(
            "PDF output requires a headless LibreOffice ('soffice') on PATH; "
            "none was found. Install LibreOffice or set VGI_DOCGEN_SOFFICE, or "
            "use DOCX output (the default)."
        )
    with tempfile.TemporaryDirectory(prefix="vgi-docgen-") as tmp:
        src = Path(tmp) / "in.docx"
        src.write_bytes(docx_bytes)
        try:
            proc = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    tmp,
                    str(src),
                ],
                capture_output=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise DocgenError(f"LibreOffice PDF conversion failed: {exc}") from exc
        out = Path(tmp) / "in.pdf"
        if proc.returncode != 0 or not out.is_file():
            detail = proc.stderr.decode("utf-8", "replace").strip() or "no PDF produced"
            raise DocgenError(f"LibreOffice PDF conversion failed: {detail}")
        return out.read_bytes()
